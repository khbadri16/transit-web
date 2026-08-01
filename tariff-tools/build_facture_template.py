"""Turn the FCT SNAX invoice into a docxtpl template (FACTURE TRANSIT.docx)."""
from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

SRC = "/sessions/clever-determined-wozniak/mnt/TRANSIT/FCT SNAX.docx"
DST = "/sessions/clever-determined-wozniak/mnt/TRANSIT/Web/frontend/api/templates/FACTURE TRANSIT.docx"

d = Document(SRC)


def set_text(p, text):
    """Replace a paragraph's text, keeping the first run's formatting."""
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        p.add_run(text)


def rebuild_cell(cell, src_p_el, texts):
    """Replace a cell's paragraphs with len(texts) clones of src_p_el."""
    tc = cell._tc
    for p in tc.findall(qn("w:p")):
        tc.remove(p)
    for t in texts:
        el = deepcopy(src_p_el)
        tc.append(el)
        set_text(Paragraph(el, cell), t)


# ---------------------------------------------------------------- header ---
P = d.paragraphs
set_text(P[0], "{{ transit_name }}")
set_text(P[2], "AGREMENT N°{{ agrement }}  NIF N°:{{ nif_transit }}")
set_text(P[5], "DOIT :  {{ client_name }}")
set_text(P[7], "Adresse : {{ client_adresse }}")
set_text(P[9], "NIF   :  {{ client_nif }}.")
set_text(P[11], "D10 N° :  {{ d10_numero }}          DU {{ d10_date }}.")

# ------------------------------------------------- table 0 (marchandise) ---
t0 = d.tables[0]
r1 = t0.rows[1]
set_text(r1.cells[0].paragraphs[1], "{{ l.marque }}")
set_text(r1.cells[1].paragraphs[1], "{{ l.designation }}")
set_text(r1.cells[2].paragraphs[1], "{{ l.nbrs }}")
set_text(r1.cells[3].paragraphs[1], "{{ l.poids }}KG")
# Repeat the row once per merchandise line. docxtpl deletes any row that
# holds a {%tr %} tag, so the loop markers get their own throw-away rows
# above and below the real one.
def tag_row(model_tr, tag):
    tr = deepcopy(model_tr)
    for tc in tr.findall(qn("w:tc")):
        for p in tc.findall(qn("w:p")):
            tc.remove(p)
        tc.append(deepcopy(model_tr.findall(qn("w:tc"))[0].findall(qn("w:p"))[0]))
    first_cell_p = tr.findall(qn("w:tc"))[0].findall(qn("w:p"))[0]
    set_text(Paragraph(first_cell_p, None), tag)
    return tr

tbl0 = t0._tbl
tbl0.insert(list(tbl0).index(r1._tr), tag_row(r1._tr, "{%tr for l in lignes %}"))
tbl0.insert(list(tbl0).index(r1._tr) + 1, tag_row(r1._tr, "{%tr endfor %}"))

# ----------------------------------------------------- table 1 (facture) ---
t1 = d.tables[1]
KEYS = [
    "echange", "magazinage", "dechargement", "restitution",
    "surestaries", "ouverture", "honnoraire", "tva", "timbre",
]
row = t1.rows[1]
src = row.cells[1].paragraphs[6]._p          # non-bold 11pt Calibri, left

rebuild_cell(row.cells[1], src, [""] + [f"{{{{ deb_{k} }}}}" for k in KEYS])
for i, k in enumerate(KEYS, start=1):        # PRODUIT already has 10 paragraphs
    set_text(row.cells[2].paragraphs[i], f"{{{{ pro_{k} }}}}")
rebuild_cell(row.cells[3], src, [""] + [f"{{{{ obs_{k} }}}}" for k in KEYS])

# TAUX row
set_text(t1.rows[2].cells[1].paragraphs[0], "{{ taux_debours }}")
set_text(t1.rows[2].cells[2].paragraphs[0], "{{ taux_produit }}")
# TOTAL TTC row — cells 1+2 are merged into one
set_text(t1.rows[3].cells[1].paragraphs[0], "{{ total_ttc }}")

d.save(DST)
print("saved:", DST)
